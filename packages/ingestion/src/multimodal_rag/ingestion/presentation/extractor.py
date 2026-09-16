"""Extract PPTX slides and convert legacy PPT files before extraction."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import time
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path


class PresentationExtractionError(RuntimeError):
    """Raised when a presentation cannot be converted or read."""


@dataclass
class PresentationOutputPaths:
    document_dir: Path
    chunks_json: Path
    metadata_json: Path
    validation_report_json: Path
    audit_markdown: Path
    tables_dir: Path | None = None


def ingest_presentation(
    presentation_path: str | Path,
    output_dir: str | Path,
    source_sha256: str | None = None,
) -> PresentationOutputPaths:
    """Extract slide text, tables, notes, and image OCR into retrieval chunks."""
    presentation_path = Path(presentation_path)
    started = time.perf_counter()
    if presentation_path.suffix.lower() == ".ppt":
        with tempfile.TemporaryDirectory(prefix="rag-ppt-") as conversion_dir:
            converted_path = _convert_ppt_to_pptx(presentation_path, Path(conversion_dir))
            slides = _read_pptx_slides(converted_path)
    else:
        slides = _read_pptx_slides(presentation_path)
    chunks = _build_chunks(presentation_path.name, slides)
    if not chunks:
        raise PresentationExtractionError("The PowerPoint file contains no extractable text.")

    document_id = f"presentation_{uuid.uuid4().hex[:12]}"
    for index, chunk in enumerate(chunks):
        chunk["metadata"]["document_id"] = document_id
        chunk["metadata"]["chunk_id"] = f"{document_id}_slide_{index:04d}"
    document_dir = Path(output_dir) / document_id
    document_dir.mkdir(parents=True, exist_ok=True)
    chunks_path = document_dir / "chunks.json"
    metadata_path = document_dir / "metadata.json"
    validation_path = document_dir / "validation_report.json"
    audit_path = document_dir / "extracted_text_audit.md"
    chunks_path.write_text(json.dumps(chunks, indent=2), encoding="utf-8")
    metadata_path.write_text(json.dumps({
        "document_id": document_id, "source_file": presentation_path.name,
        "presentation_format": presentation_path.suffix.lower().lstrip("."),
        "source_sha256": source_sha256, "slide_count": len(slides),
        "total_chunks": len(chunks), "processing_time_seconds": round(time.perf_counter() - started, 3),
    }, indent=2), encoding="utf-8")
    validation_path.write_text(json.dumps({"status": "ok", "extractor": "python-pptx", "slides": len(slides)}, indent=2), encoding="utf-8")
    audit_path.write_text("# Presentation Extraction: " + presentation_path.name + "\n\n" + "\n\n".join(chunk["chunk_text"] for chunk in chunks), encoding="utf-8")
    return PresentationOutputPaths(document_dir, chunks_path, metadata_path, validation_path, audit_path)


def _convert_ppt_to_pptx(ppt_path: Path, conversion_dir: Path) -> Path:
    soffice = _find_soffice()
    profile_uri = (conversion_dir / "libreoffice-profile").resolve().as_uri()
    try:
        completed = subprocess.run(
            [str(soffice), f"-env:UserInstallation={profile_uri}", "--headless", "--convert-to", "pptx", "--outdir", str(conversion_dir), str(ppt_path)],
            capture_output=True, text=True, timeout=120, check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise PresentationExtractionError(f"LibreOffice could not convert the PPT file: {exc}") from exc
    converted_path = conversion_dir / f"{ppt_path.stem}.pptx"
    if completed.returncode != 0 or not converted_path.exists():
        detail = (completed.stderr or completed.stdout).strip()[:500]
        raise PresentationExtractionError(f"LibreOffice could not convert the PPT file. {detail}")
    return converted_path


def _find_soffice() -> Path:
    configured_path = os.getenv("RAG_LIBREOFFICE_PATH")
    candidates = [configured_path, shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise PresentationExtractionError("LibreOffice is required for legacy .ppt uploads. Set RAG_LIBREOFFICE_PATH to soffice.exe.")


def _read_pptx_slides(pptx_path: Path) -> list[dict]:
    try:
        from pptx import Presentation
        presentation = Presentation(str(pptx_path))
    except Exception as exc:
        raise PresentationExtractionError(f"Could not read the PPTX file: {exc}") from exc
    slides: list[dict] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        ocr_used = False
        title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text.strip() else ""
        if title:
            blocks.append(title)
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        blocks.append(" | ".join(cells))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    blocks.append(text)
            elif getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                image_text = _ocr_picture(shape)
                if image_text:
                    blocks.append("Image text: " + image_text)
                    ocr_used = True
        notes = getattr(slide, "notes_slide", None)
        if notes:
            for shape in notes.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text and "click to add notes" not in text.lower():
                        blocks.append("Speaker notes: " + text)
        if blocks:
            slides.append({"slide_number": slide_number, "text": "\n".join(dict.fromkeys(blocks)), "title": title, "ocr_used": ocr_used})
    return slides


def _ocr_picture(shape) -> str:
    """Best-effort OCR for screenshots and image-only slide content."""
    try:
        from multimodal_rag.ingestion.extractors.ocr_extractor import run_ocr
        with tempfile.NamedTemporaryFile(suffix="." + shape.image.ext, delete=False) as image_file:
            image_file.write(shape.image.blob)
            image_path = Path(image_file.name)
        try:
            return run_ocr(image_path).text.strip()
        finally:
            image_path.unlink(missing_ok=True)
    except Exception:
        return ""


def _build_chunks(source_file: str, slides: list[dict]) -> list[dict]:
    timestamp = datetime.now(timezone.utc).isoformat()
    return [{"chunk_text": slide["text"], "metadata": {
        "chunk_id": "", "document_id": "", "source_file": source_file,
        "page_numbers": [slide["slide_number"]], "section_title": slide["title"] or None,
        "layout_type": "slide", "extraction_method": "python-pptx+rapidocr" if slide.get("ocr_used") else "python-pptx",
        "ocr_confidence": None, "validation_status": "ok", "ingestion_timestamp": timestamp,
        "pipeline_version": "presentation-v1", "source_region_ids": [],
        "slide_number": slide["slide_number"],
    }} for slide in slides]
