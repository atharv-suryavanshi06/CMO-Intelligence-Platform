from __future__ import annotations

import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document

from multimodal_rag.ingestion.word.extractor import ingest_word


class WordIngestionTests(unittest.TestCase):
    def test_docx_text_and_table_are_written_as_embedding_compatible_chunks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "brief.docx"
            document = Document()
            document.add_heading("Campaign brief", level=1)
            document.add_paragraph("The campaign targets returning customers with a clear value proposition.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Metric"
            table.rows[0].cells[1].text = "Target"
            document.save(source)

            paths = ingest_word(source, root / "artifacts", "source-hash")

            chunks = json.loads(paths.chunks_json.read_text(encoding="utf-8"))
            self.assertTrue(any("Campaign brief" in chunk["chunk_text"] for chunk in chunks))
            self.assertTrue(any("Metric | Target" in chunk["chunk_text"] for chunk in chunks))
            self.assertEqual(chunks[0]["metadata"]["extraction_method"], "python-docx")

    def test_legacy_doc_is_converted_then_extracted_when_libreoffice_is_available(self) -> None:
        soffice = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
        if not soffice.exists():
            self.skipTest("LibreOffice is not installed")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_docx = root / "legacy-source.docx"
            document = Document()
            document.add_paragraph("Legacy Word content is converted before extraction.")
            document.save(source_docx)
            conversion = subprocess.run(
                [
                    str(soffice), f"-env:UserInstallation={(root / 'profile').resolve().as_uri()}", "--headless",
                    "--convert-to", "doc", "--outdir", str(root), str(source_docx),
                ],
                capture_output=True, text=True, timeout=120, check=False,
            )
            source_doc = root / "legacy-source.doc"
            self.assertEqual(conversion.returncode, 0, conversion.stderr)
            self.assertTrue(source_doc.exists())

            paths = ingest_word(source_doc, root / "artifacts")

            chunks = json.loads(paths.chunks_json.read_text(encoding="utf-8"))
            self.assertIn("Legacy Word content", chunks[0]["chunk_text"])
