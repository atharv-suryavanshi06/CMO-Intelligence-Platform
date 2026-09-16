from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from multimodal_rag.ingestion.presentation.extractor import ingest_presentation


class PresentationIngestionTests(unittest.TestCase):
    def test_pptx_slide_text_table_and_metadata_are_written(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "strategy.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Growth strategy"
            slide.placeholders[1].text = "Increase retention through targeted campaigns."
            table = slide.shapes.add_table(2, 2, 100, 100, 300, 100).table
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Target"
            table.cell(1, 0).text = "Retention"
            table.cell(1, 1).text = "80%"
            presentation.save(source)

            paths = ingest_presentation(source, root / "artifacts", "source-hash")

            chunks = json.loads(paths.chunks_json.read_text(encoding="utf-8"))
            self.assertIn("Growth strategy", chunks[0]["chunk_text"])
            self.assertIn("Metric | Target", chunks[0]["chunk_text"])
            self.assertEqual(chunks[0]["metadata"]["slide_number"], 1)
            self.assertEqual(chunks[0]["metadata"]["extraction_method"], "python-pptx")
